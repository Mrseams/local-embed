"""Stage 1 - extraction: turn a file into a list of Sections.

A Section is a span of raw text plus the structural metadata we can recover
from the format (page number for PDFs, heading path for Markdown/DOCX).
"""
import re
from dataclasses import dataclass
from pathlib import Path


@dataclass
class Section:
    text: str
    page: int | None = None
    heading: str | None = None


class UnsupportedFormat(ValueError):
    pass


def load(path: Path) -> list[Section]:
    ext = path.suffix.lower()
    if ext == ".txt":
        return [Section(_read_text(path))]
    if ext in (".md", ".markdown"):
        return load_markdown(_read_text(path))
    if ext == ".pdf":
        return load_pdf(path)
    if ext == ".docx":
        return load_docx(path)
    raise UnsupportedFormat(f"Unsupported file type: {ext}")


def _read_text(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1")


_MD_HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")
_MD_FENCE = re.compile(r"^\s*(```|~~~)")


def load_markdown(text: str) -> list[Section]:
    """Split on headings, tracking the heading path (e.g. "Setup > Docker")."""
    sections: list[Section] = []
    stack: list[tuple[int, str]] = []
    buf: list[str] = []
    in_fence = False

    def flush():
        body = "\n".join(buf).strip()
        if body:
            heading = " > ".join(h for _, h in stack) or None
            sections.append(Section(body, heading=heading))
        buf.clear()

    for line in text.splitlines():
        if _MD_FENCE.match(line):
            in_fence = not in_fence
        m = None if in_fence else _MD_HEADING.match(line)
        if m:
            flush()
            level = len(m.group(1))
            while stack and stack[-1][0] >= level:
                stack.pop()
            stack.append((level, m.group(2)))
        else:
            buf.append(line)
    flush()
    return sections


def load_pdf(path: Path) -> list[Section]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as e:  # noqa: BLE001
            raise ValueError("PDF is password-protected") from e
    sections = []
    for i, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            sections.append(Section(text, page=i))
    if not sections:
        raise ValueError("No extractable text in PDF (scanned image? OCR is not supported)")
    return sections


def load_docx(path: Path) -> list[Section]:
    import docx

    document = docx.Document(str(path))
    sections: list[Section] = []
    stack: list[tuple[int, str]] = []
    buf: list[str] = []

    def flush():
        body = "\n\n".join(buf).strip()
        if body:
            heading = " > ".join(h for _, h in stack) or None
            sections.append(Section(body, heading=heading))
        buf.clear()

    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name if para.style is not None else "") or ""
        level = _docx_heading_level(style)
        if level:
            flush()
            while stack and stack[-1][0] >= level:
                stack.pop()
            stack.append((level, text))
        else:
            buf.append(text)
    flush()

    # Tables carry a lot of content in real docs; render them as pipe rows.
    for table in document.tables:
        rows = [" | ".join(c.text.strip() for c in row.cells) for row in table.rows]
        body = "\n".join(r for r in rows if r.strip(" |"))
        if body:
            sections.append(Section(body, heading="Table"))
    return sections


def _docx_heading_level(style: str) -> int:
    if style == "Title":
        return 1
    m = re.match(r"Heading (\d)", style)
    return int(m.group(1)) if m else 0
